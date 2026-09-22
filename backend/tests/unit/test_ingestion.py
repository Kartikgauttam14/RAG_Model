import json
from io import BytesIO

import pytest
from app.ingestion import ExtractionError, StructureAwareChunker, extract_document, remove_page_folios
from app.ingestion.types import ExtractedDocument, ExtractedElement
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _active_sheet(workbook: Workbook) -> Worksheet:
    """``Workbook.active`` is typed as Optional; these tests always write to the first sheet."""
    sheet = workbook.active
    assert sheet is not None
    return sheet


def test_xlsx_extraction_preserves_sheet_and_row_provenance() -> None:
    workbook = Workbook()
    sheet = _active_sheet(workbook)
    sheet.title = "Products"
    sheet.append(["SKU", "Name", "Name"])
    sheet.append(["ABM001", "Oud", "عود"])
    payload = BytesIO()
    workbook.save(payload)

    extracted = extract_document(payload.getvalue(), "catalog.xlsx", XLSX_MEDIA_TYPE)

    assert extracted.metadata["sheet_names"] == ["Products"]
    assert len(extracted.elements) == 1
    element = extracted.elements[0]
    assert "SKU: ABM001" in element.text
    assert "Name_2: عود" in element.text
    assert element.metadata == {
        "sheet": "Products",
        "sheet_index": 1,
        "row_number": 2,
        "source_type": "spreadsheet_row",
    }


def test_xlsx_extraction_skips_excluded_sheets_case_insensitively() -> None:
    workbook = Workbook()
    products = _active_sheet(workbook)
    products.title = "Products"
    products.append(["SKU", "Name"])
    products.append(["ABM001", "Oud"])
    changelog = workbook.create_sheet("01_Version_Log")
    changelog.append(["Version", "Summary"])
    changelog.append(["v2.3.8", "Loader now reads all 20 sheets"])
    payload = BytesIO()
    workbook.save(payload)

    extracted = extract_document(payload.getvalue(), "catalog.xlsx", XLSX_MEDIA_TYPE, ["01_version_log"])

    assert extracted.metadata["sheet_names"] == ["Products", "01_Version_Log"]
    assert extracted.metadata["excluded_sheets"] == ["01_Version_Log"]
    assert [element.metadata["sheet"] for element in extracted.elements] == ["Products"]
    assert all("loader now reads" not in element.text.lower() for element in extracted.elements)


def test_xlsx_extraction_keeps_every_sheet_when_no_exclusions_configured() -> None:
    workbook = Workbook()
    sheet = _active_sheet(workbook)
    sheet.title = "Products"
    sheet.append(["SKU", "Name"])
    sheet.append(["ABM001", "Oud"])
    changelog = workbook.create_sheet("01_Version_Log")
    changelog.append(["Version", "Summary"])
    changelog.append(["v2.3.8", "Loader now reads all 20 sheets"])
    payload = BytesIO()
    workbook.save(payload)

    extracted = extract_document(payload.getvalue(), "catalog.xlsx", XLSX_MEDIA_TYPE)

    assert extracted.metadata["excluded_sheets"] == []
    assert {element.metadata["sheet"] for element in extracted.elements} == {"Products", "01_Version_Log"}


def test_xlsx_extraction_drops_unpublished_workflow_rows() -> None:
    workbook = Workbook()
    sheet = _active_sheet(workbook)
    sheet.title = "Greetings"
    sheet.append(["Intent", "Text", "Status"])
    sheet.append(["greeting", "Welcome to Mansam", "Open"])
    sheet.append(["greeting", "Draft welcome line", "Claude Draft — pending MH review"])
    payload = BytesIO()
    workbook.save(payload)

    extracted = extract_document(
        payload.getvalue(),
        "catalog.xlsx",
        XLSX_MEDIA_TYPE,
        None,
        ["status", "reviewed by"],
        ["draft", "pending"],
    )

    assert len(extracted.elements) == 1
    assert "Welcome to Mansam" in extracted.elements[0].text
    assert all("draft welcome line" not in element.text.lower() for element in extracted.elements)
    assert extracted.metadata["unpublished_rows_skipped"] == 1
    assert extracted.metadata["unpublished_row_refs"] == [
        {"sheet": "Greetings", "row_number": 3, "columns": ["Status"]}
    ]


def test_xlsx_extraction_ignores_status_columns_outside_workflow_columns() -> None:
    workbook = Workbook()
    sheet = _active_sheet(workbook)
    sheet.title = "Boutiques"
    sheet.append(["City", "Status"])
    sheet.append(["Riyadh", "Open"])
    sheet.append(["Madinah", "Pending Open"])
    payload = BytesIO()
    workbook.save(payload)

    extracted = extract_document(
        payload.getvalue(),
        "catalog.xlsx",
        XLSX_MEDIA_TYPE,
        None,
        None,
        ["draft", "pending"],
    )

    assert len(extracted.elements) == 2
    assert extracted.metadata["unpublished_rows_skipped"] == 0


def test_chunker_retains_section_and_page_provenance() -> None:
    document = ExtractedDocument(
        "policy.md",
        "text/markdown",
        [
            ExtractedElement("Delivery", "heading", page_number=2),
            ExtractedElement("Orders ship in two business days.", "paragraph", page_number=2),
            ExtractedElement("Tracking is sent by email.", "paragraph", page_number=2),
        ],
    )

    chunks = StructureAwareChunker(target_chars=200, max_chars=300, overlap_chars=20).chunk(document)

    assert len(chunks) == 1
    assert chunks[0].section == "Delivery"
    assert chunks[0].page_number == 2
    assert chunks[0].metadata["source_name"] == "policy.md"


def test_chunker_overlap_repeats_table_rows_whole() -> None:
    rows = [
        ExtractedElement(
            f"Sheet: Products\nRow: {index}\nSKU: ABM-{index:03d}\nName: Perfume {index}",
            "table",
            section="Products",
        )
        for index in range(1, 13)
    ]
    document = ExtractedDocument(
        "catalog.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        rows,
    )

    chunks = StructureAwareChunker(target_chars=120, max_chars=240, overlap_chars=60).chunk(document)

    assert len(chunks) > 1
    for chunk in chunks:
        for block in chunk.content.split("\n\n"):
            assert block.startswith("Sheet: Products"), f"chunk {chunk.chunk_index} starts mid-record: {block[:40]!r}"
    emitted = [chunk.content for chunk in chunks]
    assert any("ABM-001" in content for content in emitted)
    assert any("ABM-012" in content for content in emitted)


def test_chunker_overlap_falls_back_to_character_tail_for_prose() -> None:
    document = ExtractedDocument(
        "policy.md",
        "text/markdown",
        [ExtractedElement("Orders ship in two business days and tracking follows by email.", "paragraph")],
    )

    chunks = StructureAwareChunker(target_chars=200, max_chars=300, overlap_chars=20).chunk(document)

    assert len(chunks) == 1
    assert chunks[0].content == "Orders ship in two business days and tracking follows by email."


def test_json_extraction_flattens_records_into_readable_lines() -> None:
    payload = json.dumps(
        [
            {
                "productLineId": 1,
                "productLineNameEn": "Eau de Parfum 100ml",
                "descriptionEn": "An exquisite Eau De Parfum collection with Arabian roots.",
                "image1": "Family1_EDP.jpg",
                "descriptionAr": None,
                "active": True,
                "notes": ["oud", "amber"],
                "details": {"family": "Oriental"},
            },
            {"productLineNameEn": "Luban", "descriptionEn": None, "active": False},
            {"descriptionEn": None, "active": True},
        ]
    ).encode()

    extracted = extract_document(payload, "productlines.json", "application/json")

    assert extracted.metadata == {"record_count": 3}
    assert len(extracted.elements) == 2
    element = extracted.elements[0]
    assert element.kind == "table"
    assert element.metadata == {"record_index": 0}
    assert "productLineId: 1" in element.text
    assert "productLineNameEn: Eau de Parfum 100ml" in element.text
    assert "notes[0]: oud" in element.text
    assert "notes[1]: amber" in element.text
    assert "details.family: Oriental" in element.text
    assert "descriptionAr" not in element.text
    assert "active" not in element.text
    assert extracted.elements[1].text == "productLineNameEn: Luban"
    assert extracted.elements[1].metadata == {"record_index": 1}


def test_remove_page_folios_drops_printed_page_numbers_at_page_edges() -> None:
    document = ExtractedDocument(
        "booklet.pdf",
        "application/pdf",
        [
            ExtractedElement("87\nboutiques and arts\nof mansam", "paragraph", page_number=41),
            ExtractedElement("92 93\nboutiques and arts\nof mansam", "paragraph", page_number=42),
        ],
    )

    cleaned = remove_page_folios(document)

    assert cleaned.metadata["removed_page_folios"] == 2
    assert [element.text for element in cleaned.elements] == [
        "boutiques and arts\nof mansam",
        "boutiques and arts\nof mansam",
    ]


def test_remove_page_folios_keeps_numbers_inside_the_page_body() -> None:
    document = ExtractedDocument(
        "catalog.pdf",
        "application/pdf",
        [
            ExtractedElement(
                "Boutique count\nRiyadh\n120\nJeddah\nDammam\nMecca",
                "table",
                page_number=3,
            )
        ],
    )

    cleaned = remove_page_folios(document)

    assert cleaned.metadata == {}
    assert "120" in cleaned.elements[0].text


def test_remove_page_folios_ignores_documents_without_page_numbers() -> None:
    document = ExtractedDocument(
        "sheet.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        [ExtractedElement("87\nboutiques and arts", "table")],
    )

    cleaned = remove_page_folios(document)

    assert cleaned is document
    assert cleaned.elements[0].text == "87\nboutiques and arts"


def test_json_extraction_rejects_records_without_readable_text() -> None:
    with pytest.raises(ExtractionError, match="no usable text"):
        extract_document(b'[{"descriptionEn": null, "active": true}]', "empty.json", "application/json")


def test_html_extraction_keeps_headings_and_tables() -> None:
    html = b"""
    <html><head><title>About</title></head><body>
      <h1>Delivery</h1>
      <p>Orders ship in two business days.</p>
      <table><tr><th>Region</th><th>Days</th></tr><tr><td>UAE</td><td>2</td></tr></table>
    </body></html>
    """

    extracted = extract_document(html, "about.html", "text/html")

    assert extracted.metadata["title"] == "About"
    assert extracted.elements[0].text == "Delivery"
    assert extracted.elements[0].kind == "heading"
    assert extracted.elements[0].section == "Delivery"
    assert extracted.elements[-1].kind == "table"
    assert "UAE | 2" in extracted.elements[-1].text


def test_html_extraction_reports_spa_shell_with_actionable_error() -> None:
    html = b"""
    <html><body><app-root></app-root>
      <script src="main.38242ade25b2d820.js" type="module"></script>
    </body></html>
    """

    with pytest.raises(ExtractionError, match="single-page application"):
        extract_document(html, "home.html", "text/html")
