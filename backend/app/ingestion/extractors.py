import csv
import io
import json
import re
from collections.abc import Callable
from typing import Any

from bs4 import BeautifulSoup
from bs4.element import Tag
from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader

from app.ingestion.normalize import normalize_text, remove_page_folios, remove_repeated_page_furniture
from app.ingestion.sites import SPA_EXTRACTION_HINT, is_unrendered_spa
from app.ingestion.types import ElementKind, ExtractedDocument, ExtractedElement

MARKDOWN_HEADING = re.compile(r"^(#{1,6})\s+(.+)$")
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class ExtractionError(ValueError):
    pass


def extract_document(
    data: bytes,
    name: str,
    media_type: str,
    excluded_sheets: list[str] | None = None,
    workflow_columns: list[str] | None = None,
    unpublished_markers: list[str] | None = None,
) -> ExtractedDocument:
    # ``Callable[..., ExtractedDocument]`` because the spreadsheet extractor also takes the
    # ingestion filters, while the other extractors share the three-argument shape. Without
    # the annotation mypy infers a union of signatures and refuses to call the result.
    extractors: dict[str, Callable[..., ExtractedDocument]] = {
        "application/pdf": _extract_pdf,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": _extract_docx,
        XLSX_MEDIA_TYPE: _extract_xlsx,
        "text/plain": _extract_text,
        "text/markdown": _extract_markdown,
        "text/html": _extract_html,
        "text/csv": _extract_csv,
        "application/json": _extract_json,
    }
    extractor = extractors.get(media_type)
    if extractor is None:
        raise ExtractionError(f"No extractor for {media_type}")
    try:
        if media_type == XLSX_MEDIA_TYPE:
            document = _extract_xlsx(
                data, name, media_type, excluded_sheets, workflow_columns, unpublished_markers
            )
        else:
            document = extractor(data, name, media_type)
    except Exception as exc:
        if isinstance(exc, ExtractionError):
            raise
        raise ExtractionError(f"Could not extract {name}") from exc
    if not document.elements and not document.needs_ocr:
        raise ExtractionError("Document contains no usable text")
    return document


def _extract_pdf(data: bytes, name: str, media_type: str) -> ExtractedDocument:
    reader = PdfReader(io.BytesIO(data), strict=False)
    if reader.is_encrypted:
        raise ExtractionError("Encrypted PDFs are not supported")
    elements: list[ExtractedElement] = []
    low_text_pages = 0
    for page_number, page in enumerate(reader.pages, 1):
        text = normalize_text(page.extract_text() or "")
        if len(text) < 20:
            low_text_pages += 1
            continue
        elements.append(ExtractedElement(text, "paragraph", page_number=page_number))
    needs_ocr = bool(reader.pages) and low_text_pages / len(reader.pages) > 0.5
    document = ExtractedDocument(
        name,
        media_type,
        elements,
        {
            "page_count": len(reader.pages),
            "low_text_pages": low_text_pages,
            "pdf_metadata": {str(k): str(v) for k, v in (reader.metadata or {}).items()},
        },
        needs_ocr=needs_ocr,
    )
    return remove_repeated_page_furniture(remove_page_folios(document))


def _extract_docx(data: bytes, name: str, media_type: str) -> ExtractedDocument:
    doc = DocxDocument(io.BytesIO(data))
    elements: list[ExtractedElement] = []
    section: str | None = None
    for paragraph in doc.paragraphs:
        text = normalize_text(paragraph.text)
        if not text:
            continue
        style = paragraph.style.name.lower() if paragraph.style else ""
        kind: ElementKind = "heading" if style.startswith("heading") else "paragraph"
        if kind == "heading":
            section = text
        elements.append(ExtractedElement(text, kind, section=section, metadata={"style": style}))
    for table_index, table in enumerate(doc.tables):
        rows = [[normalize_text(cell.text) for cell in row.cells] for row in table.rows]
        text = "\n".join(" | ".join(row) for row in rows if any(row))
        if text:
            elements.append(ExtractedElement(text, "table", section=section, metadata={"table_index": table_index}))
    return ExtractedDocument(name, media_type, elements, {"table_count": len(doc.tables)})


def _extract_xlsx(
    data: bytes,
    name: str,
    media_type: str,
    excluded_sheets: list[str] | None = None,
    workflow_columns: list[str] | None = None,
    unpublished_markers: list[str] | None = None,
) -> ExtractedDocument:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    elements: list[ExtractedElement] = []
    sheet_names = list(workbook.sheetnames)
    excluded = {sheet.strip().casefold() for sheet in excluded_sheets or [] if sheet.strip()}
    workflow = {column.strip().casefold() for column in workflow_columns or [] if column.strip()}
    markers = [marker.strip().casefold() for marker in unpublished_markers or [] if marker.strip()]
    skipped_sheets: list[str] = []
    unpublished_rows: list[dict[str, Any]] = []
    try:
        for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
            if sheet.title.strip().casefold() in excluded:
                skipped_sheets.append(sheet.title)
                continue
            rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
            if not rows:
                continue
            header_index = _find_header_row(rows)
            headers = _unique_headers(rows[header_index])
            for row_number, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
                if not any(value not in (None, "") for value in row):
                    continue
                fields: list[str] = []
                unpublished_in: list[str] = []
                for column_index, value in enumerate(row):
                    if value in (None, ""):
                        continue
                    header = headers[column_index] if column_index < len(headers) else f"column_{column_index + 1}"
                    fields.append(f"{header}: {_cell_text(value)}")
                    if header.strip().casefold() in workflow and _matches_marker(value, markers):
                        unpublished_in.append(header)
                if unpublished_in:
                    unpublished_rows.append(
                        {
                            "sheet": sheet.title,
                            "row_number": row_number,
                            "columns": unpublished_in,
                        }
                    )
                    continue
                elements.append(
                    ExtractedElement(
                        text=f"Sheet: {sheet.title}\nRow: {row_number}\n" + "\n".join(fields),
                        kind="table",
                        section=sheet.title,
                        metadata={
                            "sheet": sheet.title,
                            "sheet_index": sheet_index,
                            "row_number": row_number,
                            "source_type": "spreadsheet_row",
                        },
                    )
                )
    finally:
        workbook.close()
    return ExtractedDocument(
        name,
        media_type,
        elements,
        {
            "sheet_count": len(sheet_names),
            "sheet_names": sheet_names,
            "excluded_sheets": skipped_sheets,
            "unpublished_rows_skipped": len(unpublished_rows),
            "unpublished_row_refs": unpublished_rows,
        },
    )


def _find_header_row(rows: list[tuple[Any, ...]]) -> int:
    best_index = 0
    best_score = -1
    for index, row in enumerate(rows[:10]):
        nonempty = [value for value in row if value not in (None, "")]
        strings = [value for value in nonempty if isinstance(value, str)]
        score = (len(strings) * 2) + len(nonempty)
        if len(nonempty) >= 2 and score > best_score:
            best_index = index
            best_score = score
    return best_index


def _unique_headers(row: tuple[Any, ...]) -> list[str]:
    seen: dict[str, int] = {}
    headers: list[str] = []
    for index, value in enumerate(row, start=1):
        base = str(value).strip() if value not in (None, "") else f"column_{index}"
        seen[base] = seen.get(base, 0) + 1
        headers.append(base if seen[base] == 1 else f"{base}_{seen[base]}")
    return headers


def _cell_text(value: Any) -> str:
    if isinstance(value, str):
        return " ".join(value.split())
    return str(value)


def _matches_marker(value: Any, markers: list[str]) -> bool:
    """Report whether a workflow cell holds an unpublished/authoring marker such as ``Draft``.

    Markers are matched as whole words so that legitimate values are never dropped by
    accident (for example ``Status: "Depending on stock"`` must not match ``pending``).
    """
    if not markers:
        return False
    text = _cell_text(value).casefold()
    return any(
        re.search(rf"(?<![a-z0-9]){re.escape(marker)}(?![a-z0-9])", text) is not None for marker in markers
    )


def _extract_text(data: bytes, name: str, media_type: str) -> ExtractedDocument:
    text = _decode_text(data)
    paragraphs = [normalize_text(part) for part in re.split(r"\n\s*\n", text)]
    elements = [ExtractedElement(part, "paragraph") for part in paragraphs if part]
    return ExtractedDocument(name, media_type, elements)


def _extract_markdown(data: bytes, name: str, media_type: str) -> ExtractedDocument:
    text = _decode_text(data)
    elements: list[ExtractedElement] = []
    section: str | None = None
    buffer: list[str] = []
    in_code = False

    def flush(kind: ElementKind = "paragraph") -> None:
        if buffer:
            content = normalize_text("\n".join(buffer))
            if content:
                elements.append(ExtractedElement(content, kind, section=section))
            buffer.clear()

    for line in text.splitlines():
        if line.strip().startswith("```"):
            if in_code:
                flush("code")
            else:
                flush()
            in_code = not in_code
            continue
        match = MARKDOWN_HEADING.match(line)
        if match and not in_code:
            flush()
            section = normalize_text(match.group(2))
            elements.append(
                ExtractedElement(section, "heading", section=section, metadata={"level": len(match.group(1))})
            )
        elif not line.strip() and not in_code:
            flush()
        else:
            buffer.append(line)
    flush("code" if in_code else "paragraph")
    return ExtractedDocument(name, media_type, elements)


def _extract_html(data: bytes, name: str, media_type: str) -> ExtractedDocument:
    raw_text = _decode_text(data)
    soup = BeautifulSoup(raw_text, "html.parser")
    for tag in soup(["script", "style", "noscript", "template", "svg"]):
        tag.decompose()
    elements: list[ExtractedElement] = []
    section: str | None = None
    for raw_node in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "pre", "table"]):
        if not isinstance(raw_node, Tag):
            continue
        node = raw_node
        node_name = str(node.name)
        text = normalize_text(node.get_text(" ", strip=True))
        if not text:
            continue
        kind: ElementKind
        if node_name.startswith("h"):
            section = text
            kind = "heading"
        elif node_name == "table":
            rows = []
            for raw_row in node.find_all("tr"):
                if not isinstance(raw_row, Tag):
                    continue
                cells = [
                    normalize_text(cell.get_text(" ", strip=True))
                    for raw_cell in raw_row.find_all(["th", "td"])
                    if isinstance(raw_cell, Tag)
                    for cell in [raw_cell]
                ]
                rows.append(" | ".join(cells))
            text = "\n".join(row for row in rows if row)
            kind = "table"
        elif node_name == "li":
            kind = "list"
        elif node_name == "pre":
            kind = "code"
        else:
            kind = "paragraph"
        elements.append(ExtractedElement(text, kind, section=section))
    if not elements and is_unrendered_spa(raw_text):
        raise ExtractionError(SPA_EXTRACTION_HINT)
    return ExtractedDocument(name, media_type, elements, {"title": soup.title.string if soup.title else None})


def _extract_csv(data: bytes, name: str, media_type: str) -> ExtractedDocument:
    reader = csv.DictReader(io.StringIO(_decode_text(data)))
    elements = []
    for row_number, row in enumerate(reader, 2):
        text = "\n".join(f"{key}: {value}" for key, value in row.items() if value not in (None, ""))
        if text:
            elements.append(ExtractedElement(text, "table", metadata={"row_number": row_number}))
    return ExtractedDocument(name, media_type, elements, {"columns": reader.fieldnames or []})


def _extract_json(data: bytes, name: str, media_type: str) -> ExtractedDocument:
    payload: Any = json.loads(_decode_text(data))
    records = payload if isinstance(payload, list) else [payload]
    elements = []
    for index, record in enumerate(records):
        text = _render_json_record(record)
        if not text:
            continue
        elements.append(ExtractedElement(text, "table", metadata={"record_index": index}))
    return ExtractedDocument(name, media_type, elements, {"record_count": len(records)})


def _render_json_record(record: Any) -> str:
    lines: list[str] = []
    _collect_json_lines(record, "", lines)
    return "\n".join(lines)


def _collect_json_lines(value: Any, prefix: str, lines: list[str]) -> None:
    """Flatten nested JSON into ``path: value`` lines, skipping empty and boolean fields."""
    if isinstance(value, dict):
        for key in sorted(value, key=str):
            _collect_json_lines(value[key], f"{prefix}{key}.", lines)
        return
    if isinstance(value, list):
        for index, item in enumerate(value):
            _collect_json_lines(item, f"{prefix.rstrip('.')}[{index}].", lines)
        return
    if value is None or isinstance(value, bool):
        return
    text = normalize_text(str(value))
    if not text:
        return
    lines.append(f"{prefix.rstrip('.') or 'value'}: {text}")


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "utf-16", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("Text encoding could not be determined")
